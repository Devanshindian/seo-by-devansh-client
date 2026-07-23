#!/usr/bin/env python3
"""Step H — build the two clean end-files. The last step; it invents nothing (pure assembly, F2).

Reads:  output/format-summary.csv   (Step F — which shapes earn links)
        _work/ideas-ranked.json     (Step G3 — the ranked distinct ideas)
        _work/g2_sheet.json          (the master — every kept page, its asset/gap/angle/idea)
Writes: output/competitor-study-ideas.xlsx   THE DELIVERABLE (decision file): Tab 1 formats, Tab 2 ideas
        output/competitor-formats.xlsx        the MASTER / evidence file (one row per kept page)

Two files, never the same data twice (recipe H). The ideas tab is coloured by brand fit (CORE green /
TRANSPLANT blue / ADJACENT amber) and carries the Backing URLs inline, so the decision file stands alone.
"""
import argparse, csv, json, os, sys
import config as c
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

_FILL = {"CORE": "C6EFCE", "TRANSPLANT": "BDD7EE", "ADJACENT": "FFE699"}   # green / blue / amber
_HEAD = PatternFill("solid", fgColor="404040")
_HEADFONT = Font(bold=True, color="FFFFFF")


def _autosize(ws, maxw=70):
    for col in ws.columns:
        w = max((len(str(cell.value)) for cell in col if cell.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(w + 2, maxw)


def _header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEAD; cell.font = _HEADFONT; cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"


def run(redo=False):
    fmt_csv = os.path.join(c.OUT, "format-summary.csv")
    ideas_json = os.path.join(c.WORK, "ideas-ranked.json")
    if not os.path.exists(ideas_json):
        sys.exit("!! no _work/ideas-ranked.json — run step_g3_ideas.py first")

    # ---- FILE 1: the decision file --------------------------------------------------------------
    wb = Workbook()
    # Tab 1 — Format summary (Step F), lifted verbatim
    ws1 = wb.active; ws1.title = "Format summary"
    if os.path.exists(fmt_csv):
        rows = list(csv.reader(open(fmt_csv)))
        _header(ws1, rows[0])
        for r in rows[1:]:
            ws1.append(r)
    else:
        _header(ws1, ["(format-summary.csv not found — run step_f_formats.py)"])
    _autosize(ws1)

    # Tab 2 — Ideas (Step G3), ranked, coloured by brand fit
    ideas = json.load(open(ideas_json))
    ws2 = wb.create_sheet("Ideas")
    cols = ["rank", "build_window", "brand_fit", "asset", "format", "tool_escalation", "distinct_angle",
            "total_domains", "competitors", "backing_pages", "median_words", "backing_urls"]
    _header(ws2, ["#", "Build window", "Brand fit", "Asset", "Format", "Tool escalation", "Distinct angle",
                  "Total domains", "# comps", "# pages", "Median words (typical page len)", "Backing URLs"])
    for it in ideas:
        ws2.append([it.get(k, "") for k in cols])
        fill = _FILL.get(it.get("brand_fit", ""))
        if fill:
            for cell in ws2[ws2.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
    for col_letter in ("D", "G"):        # wrap the long text columns (Asset, Distinct angle)
        for cell in ws2[col_letter]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws2)
    ws2.column_dimensions["D"].width = 55; ws2.column_dimensions["F"].width = 24   # Asset · Tool escalation
    ws2.column_dimensions["G"].width = 50; ws2.column_dimensions["K"].width = 18   # Distinct angle · Median words
    p1 = os.path.join(c.OUT, "competitor-study-ideas.xlsx")
    wb.save(p1 + ".tmp"); os.replace(p1 + ".tmp", p1)

    # ---- FILE 2: the master / evidence file -----------------------------------------------------
    master = json.load(open(os.path.join(c.WORK, "g2_sheet.json")))
    wb2 = Workbook(); wsm = wb2.active; wsm.title = "Master"
    mcols = ["competitor", "url", "format", "domains_follow", "backlinks", "asset", "brand_fit",
             "angle_gap", "distinct_angle", "tool_escalation", "g_notes"]
    _header(wsm, ["Competitor", "URL", "Format", "Follow domains", "Backlinks", "Asset", "Brand fit",
                  "Angle gap", "Distinct angle", "Tool escalation", "Notes"])
    for r in master:
        wsm.append([r.get(k, "") for k in mcols])
    _autosize(wsm, maxw=45)
    p2 = os.path.join(c.OUT, "competitor-formats.xlsx")
    wb2.save(p2 + ".tmp"); os.replace(p2 + ".tmp", p2)

    print(f"   DECISION FILE: {c.rel(p1)}")
    print(f"     Tab 1 'Format summary' · Tab 2 'Ideas' — {len(ideas)} ranked ideas, coloured by brand fit")
    print(f"   MASTER FILE:   {c.rel(p2)}  ({len(master)} rows of evidence)")
    return p1, p2


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--redo", action="store_true")
    run(ap.parse_args().redo)
